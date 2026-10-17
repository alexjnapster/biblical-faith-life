"""Optional native PDF/PPTX exports with editable evidence and explicit sources.

Base helpers remain standard-library-only. These exports require the optional
dependencies listed in native-export-requirements.txt. No package installation,
model call, image generation, or external request occurs here.
"""
from pathlib import Path
import hashlib
import os
import tempfile
from .visuals import render

WORDS={
 'uk':{'source':'Джерело','certainty':'Статус свідчень','coordinate':'Координати','dates':'Діапазон дат','continued':'продовження','text':'Текст','historical':'Історичні дані','inferred':'Висновок','reconstruction':'Реконструкція','uncertain':'Невизначено'},
 'en':{'source':'Source','certainty':'Evidence status','coordinate':'Coordinates','dates':'Date range','continued':'continued','text':'Text','historical':'Historical evidence','inferred':'Inference','reconstruction':'Reconstruction','uncertain':'Uncertain'},
 'ru':{'source':'Источник','certainty':'Статус свидетельств','coordinate':'Координаты','dates':'Диапазон дат','continued':'продолжение','text':'Текст','historical':'Исторические данные','inferred':'Вывод','reconstruction':'Реконструкция','uncertain':'Неопределённо'}}

def font_path(explicit=None):
    candidates=[explicit] if explicit else [
        '/System/Library/Fonts/Supplemental/Arial.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        'C:/Windows/Fonts/arial.ttf']
    for path in candidates:
        if path and Path(path).is_file():return str(Path(path).resolve())
    raise ValueError('A Unicode TrueType font is required; provide --font with an authorized local font file')

def _font(path):
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:raise RuntimeError('Native export requires optional reportlab; see native-export-requirements.txt') from None
    name='BFL-'+hashlib.sha256(Path(path).read_bytes()).hexdigest()[:16]
    if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,path))
    return name

def _wrap(text,name,size,width=864):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    result=[]
    for paragraph in text.split('\n'):
        line=''
        for word in paragraph.split():
            proposed=line+' '+word if line else word
            if stringWidth(proposed,name,size)<=width:line=proposed;continue
            if line:result.append(line);line=''
            # A long unbroken source URL must not overflow the page.
            for character in word:
                if line and stringWidth(line+character,name,size)>width:result.append(line);line=''
                line+=character
        result.append(line)
    return result

def _pages(spec,name):
    labels=WORDS[spec['language']];pages=[]
    def add(title,body):
        heading=_wrap(title,name,32)
        if len(heading)>3:
            body=title+'\n\n'+body;heading=heading[:2]+['…']
        lines=_wrap(body,name,22);available=max(1,int((444-len(heading)*38)/29))
        for offset in range(0,max(1,len(lines)),available):
            pages.append({'title':heading,'lines':lines[offset:offset+available],'diagram':None})
    if spec['kind'] in {'map','timeline'} or (spec['kind']=='reconstruction' and spec.get('elements')):
        pages.append({'title':_wrap(spec['title'],name,26)[:2],'lines':[],'diagram':spec})
    add(spec['title'],spec['description'])
    for index,item in enumerate(spec['items'],1):
        body=item['detail']+'\n\n'+labels['certainty']+': '+labels[item['certainty']]
        if spec['kind']=='map':body+='\n'+labels['coordinate']+': '+str(item['lat'])+', '+str(item['lon'])
        if spec['kind']=='timeline':body+='\n'+labels['dates']+': '+str(item['start'])+' — '+str(item['end'])
        body+='\n\n'+labels['source']+': '+item['source']
        add(str(index)+'. '+item['label'],body)
    if len(pages)>500:raise ValueError('Native export exceeds 500 pages')
    return pages

def _geometry(spec):
    """Native evidence diagrams, not decorative illustrations. Top-left coordinates."""
    shapes=[]
    if spec['kind']=='map':
        xs=[i['lon'] for i in spec['items']];ys=[i['lat'] for i in spec['items']]
        left=min(xs)-1;right=max(xs)+1;bottom=min(ys)-1;top=max(ys)+1
        for index,item in enumerate(spec['items'],1):
            x=65+(item['lon']-left)/(right-left)*820;y=130+(top-item['lat'])/(top-bottom)*270
            shapes.extend([('circle',x,y,4,4,False),('label',x+8,y-8,str(index+spec.get('_index_offset',0)),None,False)])
    elif spec['kind']=='timeline':
        items=spec['items'];start=min(i['start'] for i in items);end=max(i['end'] for i in items);span=max(1,end-start)
        shapes.extend([('line',65,130,895,130,False),('label',65,105,str(start),None,False),('label',830,105,str(end),None,False)])
        for index,item in enumerate(items,1):
            x=65+(item['start']-start)/span*830;x2=65+(item['end']-start)/span*830;y=155+(index-1)*min(26,260/max(1,len(items)))
            shapes.extend([('line',x,y,max(x+2,x2),y,False),('label',max(40,min(x+6,880)),y-17,str(index+spec.get('_index_offset',0)),None,False)])
    elif spec['kind']=='reconstruction':
        for e in spec.get('elements',[]):
            uncertain=spec['items'][e['source_item']-1]['certainty'] in {'uncertain','reconstruction'}
            x=65+e['x']*8.3;y=130+e['y']*2.7
            if e['shape']=='rect':shape=('rect',x,y,e['width']*8.3,e['height']*2.7,uncertain)
            elif e['shape']=='line':shape=('line',x,y,65+e['x2']*8.3,130+e['y2']*2.7,uncertain)
            else:shape=('circle',x,y,e['r']*2.7,e['r']*2.7,uncertain)
            shapes.extend([shape,('label',x+5,y+4,str(e['source_item']),None,False)])
    return shapes

def export(spec,destination,format,font=None):
    render(spec) # Same validation as SVG/HTML; data remain untrusted.
    if any(str(key).startswith('_') for key in spec):raise ValueError('Reserved native metadata field')
    if format not in {'pdf','pptx'}:raise ValueError('Native format must be pdf or pptx')
    destination=Path(destination)
    if destination.exists() or destination.is_symlink():raise ValueError('Native export must use a new output path')
    path=font_path(font);name=_font(path);pages=_pages(spec,name)
    # Dense overview diagrams are split instead of silently overlapping labels.
    if spec['kind'] in {'map','timeline'} and len(spec['items'])>12:
        pages=[p for p in pages if p['diagram'] is None]
        for offset in range(0,len(spec['items']),12):
            subset=dict(spec,items=spec['items'][offset:offset+12],_index_offset=offset);pages.insert(offset//12,{'title':_wrap(spec['title'],name,26)[:2],'lines':[],'diagram':subset})
    destination.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent,suffix='.'+format,delete=False) as handle:temporary=Path(handle.name)
    try:
        if format=='pdf':_pdf(spec,pages,temporary,name)
        else:_pptx(spec,pages,temporary,path)
        os.replace(temporary,destination)
    finally:
        if temporary.exists():temporary.unlink()
    return {'path':str(destination.resolve()),'format':format,'pages':len(pages),'sha256':hashlib.sha256(destination.read_bytes()).hexdigest(),'scope':'Native export created; rendering and source accuracy require separate verification'}

def _pdf(spec,pages,path,font):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.colors import HexColor
    c=Canvas(str(path),pagesize=(960,540),pageCompression=1);c.setTitle(spec['title']);c.setAuthor('biblical-faith-life')
    for page in pages:
        c.setFillColor(HexColor('#17252b'));c.setStrokeColor(HexColor('#17252b'));c.setFont(font,26 if page['diagram'] else 32)
        for i,line in enumerate(page['title']):c.drawString(48,490-i*38,line)
        c.setFont(font,22);y=490-len(page['title'])*38-16
        for line in page['lines']:c.drawString(48,y,line);y-=29
        for kind,x,y,a,b,dash in _geometry(page['diagram']) if page['diagram'] else []:
            c.setDash([6,4] if dash else [])
            if kind=='line':c.line(x,540-y,a,540-b)
            elif kind=='rect':c.rect(x,540-y-b,a,b,fill=0)
            elif kind=='circle':c.circle(x,540-y,a,fill=0)
            else:c.setFont(font,16);c.drawString(x,540-y-16,a)
        if page['diagram']:
            c.setFont(font,13)
            for i,line in enumerate(_wrap(spec['description'],font,13)[:4]):c.drawString(48,80-i*16,line)
        c.showPage()
    c.save()

def _pptx(spec,pages,path,font):
    try:
        from pptx import Presentation
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
    except ImportError:raise RuntimeError('PowerPoint export requires optional python-pptx; see native-export-requirements.txt') from None
    from reportlab.pdfbase.ttfonts import TTFont
    family=TTFont('font-probe',font).face.familyName
    if isinstance(family,bytes):family=family.decode('utf-8',errors='replace')
    presentation=Presentation();presentation.slide_width=Pt(960);presentation.slide_height=Pt(540);presentation.core_properties.title=spec['title'];presentation.core_properties.author='biblical-faith-life';presentation.core_properties.language=spec['language']
    def box(slide,text,x,y,width,height,size):
        shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(width),Pt(height));frame=shape.text_frame;frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0;frame.word_wrap=False
        for index,line in enumerate(text.split('\n')):
            paragraph=frame.paragraphs[0] if index==0 else frame.add_paragraph();paragraph.text=line;paragraph.font.name=family;paragraph.font.size=Pt(size);paragraph.font.color.rgb=RGBColor.from_string('17252B');paragraph.space_before=Pt(0);paragraph.space_after=Pt(0);paragraph.line_spacing=Pt(38 if size>=26 else 29 if size==22 else 18)
        return shape
    for page in pages:
        slide=presentation.slides.add_slide(presentation.slide_layouts[6]);box(slide,'\n'.join(page['title']),48,26,864,len(page['title'])*38+5,26 if page['diagram'] else 32)
        if page['lines']:box(slide,'\n'.join(page['lines']),48,26+len(page['title'])*38+16,864,458-len(page['title'])*38,22)
        for kind,x,y,a,b,dash in _geometry(page['diagram']) if page['diagram'] else []:
            if kind=='label':box(slide,a,x,y,80,24,16);continue
            if kind=='line':shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x),Pt(y),Pt(a),Pt(b))
            elif kind=='rect':shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(a),Pt(b));shape.fill.background()
            else:shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(x-a),Pt(y-a),Pt(a*2),Pt(a*2));shape.fill.background()
            shape.line.color.rgb=RGBColor.from_string('17252B');shape.line.width=Pt(1.5)
            if dash:shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        if page['diagram']:box(slide,'\n'.join(_wrap_for_notes(spec['description'])[:4]),48,435,864,80,13)
        note=spec['description']+'\n\n'+'\n\n'.join(item['label']+'\n'+item['detail']+'\n'+item['source']+'\n'+WORDS[spec['language']][item['certainty']] for item in spec['items'])
        slide.notes_slide.notes_text_frame.text=note
    presentation.save(str(path))

def _wrap_for_notes(text):
    # Diagrams have a concise caption; full description remains in following
    # editable pages and in every slide's speaker notes.
    import textwrap
    return textwrap.wrap(text,110)
